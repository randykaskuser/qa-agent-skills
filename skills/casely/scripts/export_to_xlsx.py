"""
Casely Export — turns generated Markdown test cases into Excel for TMS import.

By default every test case lands in ONE workbook, one row per case, because that
is what TestRail/Qase/Zephyr bulk import expects. Use --split when you want one
file per case instead (useful for per-case review rather than import).

The exporter refuses to write a file it cannot read faithfully. A test case that
silently loses its steps is worse than no export at all, so malformed input is
reported and skipped rather than patched up.
"""

import re
import sys
import argparse
from pathlib import Path
from typing import Dict, List, Tuple

try:
    from openpyxl import Workbook
    from openpyxl.worksheet.worksheet import Worksheet
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    print("Error: openpyxl is required. Install it with: pip install openpyxl")
    sys.exit(1)

MIN_COL_WIDTH = 10
MAX_COL_WIDTH = 60

DEFAULT_RESULTS_DIR = "results"
DEFAULT_OUTPUT_DIR = "exports"
DEFAULT_COMBINED_NAME = "all_test_cases.xlsx"
SHEET_TITLE = "Test Cases"


class TableError(Exception):
    """A Markdown test case that cannot be read without losing data."""


def _split_table_row(line: str) -> List[str]:
    """Split a Markdown table row into individual cell values."""
    raw_parts = re.split(r'(?<!\\)\|', line)
    start: int = 1 if line.startswith('|') else 0
    end: int = -1 if line.endswith('|') else len(raw_parts)

    result: List[str] = []
    raw_len = len(raw_parts)
    for i in range(start, int(end) if end >= 0 else raw_len + int(end)):
        cell = raw_parts[i]
        if cell is not None:
            result.append(cell.strip().replace(r'\|', '|'))
    return result


def _is_separator(line: str) -> bool:
    return bool(re.match(r'^\|[\s\-:|]+\|$', line))


def parse_md_table(md_content: str) -> Tuple[List[str], List[List[str]]]:
    """Parse a Markdown table into headers and data rows.

    Raises TableError when the table cannot be read faithfully: a raw newline
    inside a cell or an unescaped '|' both corrupt the row, and quietly padding
    or trimming would hand the user a plausible-looking but wrong test case.
    """
    lines = [line.strip() for line in md_content.strip().split('\n')]
    table_line_numbers = [i for i, line in enumerate(lines) if line.startswith('|')]

    if len(table_line_numbers) < 2:
        return [], []

    first, last = table_line_numbers[0], table_line_numbers[-1]

    # A real line break inside a cell is the most common way a generated case
    # gets mangled: Markdown ends the row at the newline, so everything after it
    # drops out of the export. Two symptoms give it away — stray text between
    # table rows, and a row that starts with '|' but never closes with one.
    newline_break = (
        [lines[i] for i in range(first, last + 1)
         if lines[i] and not lines[i].startswith('|')]
        or [lines[i] for i in table_line_numbers if not lines[i].endswith('|')]
    )
    if newline_break:
        raise TableError(
            f"the row is cut off at {newline_break[0][:60]!r}. A cell contains a "
            "real line break — write it as '<br>' so the whole case stays on one row."
        )

    headers = _split_table_row(lines[first])
    rows: List[List[str]] = []

    for i in table_line_numbers[1:]:
        line = lines[i]
        if _is_separator(line):
            continue

        row = _split_table_row(line)
        if not row:
            continue

        if len(row) != len(headers):
            raise TableError(
                f"row has {len(row)} cells but the header has {len(headers)}. "
                "An unescaped '|' inside the text splits the row — write it as "
                "'\\|' so the columns stay aligned."
            )
        rows.append(row)

    return headers, rows


def read_cases(results_path: Path) -> Tuple[List[Tuple[Path, List[str], List[List[str]]]], List[str]]:
    """Read every Markdown test case, separating good ones from broken ones."""
    cases: List[Tuple[Path, List[str], List[List[str]]]] = []
    problems: List[str] = []

    for md_file in sorted(results_path.glob('*.md')):
        try:
            headers, rows = parse_md_table(md_file.read_text(encoding='utf-8'))
        except TableError as exc:
            problems.append(f"{md_file.name}: {exc}")
            continue

        if not headers:
            problems.append(f"{md_file.name}: no Markdown table found.")
            continue
        if not rows:
            problems.append(f"{md_file.name}: table has headers but no test case row.")
            continue

        cases.append((md_file, headers, rows))

    return cases, problems


def _write_sheet(ws: Worksheet, headers: List[str], rows: List[List[str]]) -> None:
    """Write one header row plus data rows, formatted for reading."""
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for row_idx, row in enumerate(rows, 2):
        for col_idx, value in enumerate(row, 1):
            clean_value = value.replace('<br>', '\n').replace('<BR>', '\n') if value else ''
            cell = ws.cell(row=row_idx, column=col_idx, value=clean_value)
            cell.alignment = Alignment(wrap_text=True, vertical='top')

    for col_idx in range(1, len(headers) + 1):
        max_len = 0
        for row_idx in range(1, len(rows) + 2):
            cell_val = str(ws.cell(row=row_idx, column=col_idx).value or '')
            for line in cell_val.split('\n'):
                max_len = max(max_len, len(line))
        width = min(max(max_len + 2, MIN_COL_WIDTH), MAX_COL_WIDTH)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = 'A2'


def export_combined(
    cases: List[Tuple[Path, List[str], List[List[str]]]],
    out_dir: Path,
    file_name: str,
) -> Path:
    """Write every test case into one workbook, one row per case.

    Cases are grouped by their column structure. Normally the style guide makes
    that identical for all of them and the result is a single sheet; a mismatch
    gets its own sheet rather than being force-fitted into the wrong columns.
    """
    groups: Dict[Tuple[str, ...], List[List[str]]] = {}
    for _, headers, rows in cases:
        groups.setdefault(tuple(headers), []).extend(rows)

    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]

    for index, (headers, rows) in enumerate(groups.items()):
        title = SHEET_TITLE if index == 0 else f"{SHEET_TITLE} {index + 1}"
        ws = wb.create_sheet(title=title)
        _write_sheet(ws, list(headers), rows)

    if len(groups) > 1:
        print(
            f"Warning: {len(groups)} different column structures found — each is on its own "
            "sheet. Check that every test case followed the same style guide."
        )

    dest = out_dir / file_name
    wb.save(str(dest))
    return dest


def export_split(
    cases: List[Tuple[Path, List[str], List[List[str]]]],
    out_dir: Path,
) -> List[Path]:
    """Write one workbook per test case, named after its source file."""
    written: List[Path] = []
    for md_file, headers, rows in cases:
        wb = Workbook()
        ws = wb.active
        ws.title = "Test Case"  # type: ignore[union-attr]
        _write_sheet(ws, headers, rows)  # type: ignore[arg-type]

        dest = out_dir / f"{md_file.stem}.xlsx"
        wb.save(str(dest))
        written.append(dest)
    return written


def main() -> None:
    arg_parser = argparse.ArgumentParser(
        description='Casely Export — Markdown test cases to Excel'
    )
    arg_parser.add_argument(
        'results_dir', nargs='?', default=DEFAULT_RESULTS_DIR,
        help=f"Directory holding the .md test cases (default: '{DEFAULT_RESULTS_DIR}')",
    )
    arg_parser.add_argument(
        'output_path', nargs='?', default=DEFAULT_OUTPUT_DIR,
        help=f"Directory to write Excel files to (default: '{DEFAULT_OUTPUT_DIR}')",
    )
    arg_parser.add_argument(
        '--split', action='store_true',
        help='Write one Excel file per test case instead of one combined workbook',
    )
    arg_parser.add_argument(
        '--name', default=DEFAULT_COMBINED_NAME,
        help=f"Name of the combined workbook (default: '{DEFAULT_COMBINED_NAME}')",
    )
    args = arg_parser.parse_args()

    results_path = Path(args.results_dir)
    if not results_path.exists():
        print(f"Error: Results directory not found: {args.results_dir}")
        sys.exit(1)

    out_dir = Path(args.output_path)
    out_dir.mkdir(parents=True, exist_ok=True)

    cases, problems = read_cases(results_path)

    if not cases and not problems:
        print(f"Warning: No .md files found in {args.results_dir}")
        return

    if cases:
        if args.split:
            written = export_split(cases, out_dir)
            print(f"Exported {len(written)} files to {out_dir}/")
        else:
            dest = export_combined(cases, out_dir, args.name)
            total_rows = sum(len(rows) for _, _, rows in cases)
            print(f"Exported {total_rows} test cases to {dest}")

    if problems:
        print(f"\n{len(problems)} test case(s) were NOT exported:")
        for problem in problems:
            print(f"  - {problem}")
        print("\nFix the Markdown above and run the export again.")
        sys.exit(1)


if __name__ == '__main__':
    main()
